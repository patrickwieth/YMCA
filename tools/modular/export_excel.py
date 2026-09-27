"""Export editable Excel what-if workbook; JSON remains the source of truth."""
import argparse
import json
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.workbook.properties import CalcProperties
from openpyxl.utils import get_column_letter as col

from calibrate import calculate
from fit_prices import family_targets, scan, CP_CANDIDATES

FIELDS = [
    ('ID', None), ('Rolle', 'role'), ('Fraktionen', 'factions'), ('Katalog-Tier', 'tier'),
    ('CP', 'cp'), ('Tech', 'tech'), ('Hardwarepreis', 'cost'), ('Masse kg', 'mass'),
    ('Bedarf kWe', 'electric_kw'), ('Motor kWm', 'mechanical_kw'),
    ('Generator Wirkungsgrad', 'efficiency'), ('Generator max kWe', 'max_electric_kw'),
    ('HP Chassis', 'hp'), ('Panzerung', 'armor_type'), ('Chassis Masse %', 'mass_percent'),
    ('Chassis HP %', 'hp_percent'), ('Chassis Preis %', 'cost_percent'),
    ('Referenzmasse kg', 'reference_mass'), ('Referenzleistung kWm', 'reference_kw'),
    ('Referenztempo', 'reference_speed'), ('Max Tempo', 'max_speed'), ('Max Masse kg', 'max_mass'),
    ('Waffentraegerplaetze', 'carrier_slots'), ('Zusatzplaetze', 'equipment_slots'),
    ('Salve', 'burst'), ('Schussabstand Ticks', 'burst_delay_ticks'), ('Nachladen Ticks', 'reload_ticks'),
    ('Reichweite Zellen', 'range_cells'), ('Rohschaden', 'damage'), ('Hinweis', 'note'),
    ('Quelle', 'source'), ('Kompatibilitaet', 'allowed'), ('Ausrüstung erlaubt', 'equipment'),
    ('Waffen erlaubt', 'weapons'), ('Munition erlaubt', 'ammunition'),
    ('Anzeigename', 'display_name'), ('Hardware Fertigung %', 'hardware_percent'),
    ('Munition Reichweite Override', 'range_override_cells'),
    ('Transportkapazitaet', 'cargo_capacity'),
    ('Ausschluesse', 'excludes'), ('Aufsatzplatz', 'attachment_slot'),
    ('Zusatztraeger erlaubt', 'auxiliary_mount'), ('Infanterie erlaubt', 'allowed_infantry')]


def table(ws, name):
    if ws.max_row < 2:
        return
    t = Table(displayName=name, ref=f'A1:{col(ws.max_column)}{ws.max_row}')
    t.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
    ws.add_table(t)
    ws.freeze_panes = 'B2'
    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='204060')
    for column in ws.columns:
        ws.column_dimensions[column[0].column_letter].width = min(42, max(16, len(str(column[0].value)) + 2))


def build(catalog, output):
    wb = Workbook()
    info = wb.active
    info.title = 'Hinweise'
    for row in [
        ['YMCA Modular – Kalibrierung', ''],
        ['Referenz', f"Quellblock-Abgleich: {catalog['reference_commit']} (integrierter Game-Stand), siehe {catalog['reference_audit']}."],
        ['Status', 'Experimentelle Zahlen; keine Spielregeln oder grafische/volle Verhaltensgleichheit.'],
        ['Eingaben', 'Parameter und Bausteinwerte sind editierbar. Formeln berechnen Entwuerfe/Vergleich beim Öffnen neu.'],
        ['Preise', 'Hardwarekosten minus CP-Anforderung mal globaler CP-Creditwert. Keine Sonderrabatte.'],
        ['Massenproduktion', 'Separates Fertigungsmodul: vorläufig 95% aller Hardwarekosten, danach globaler CP-Abzug. Preisfit ist separate Studie.'],
        ['Arbeitsfluss', 'catalog.json ist Quelle. export_excel.py erzeugt eine neue Momentaufnahme, kein Excel-Import.'],
        ['Schutz', 'Exporter überschreibt keine bestehende Arbeitsmappe; für neue Daten neuen Dateinamen verwenden.'],
        ['Erweiterung', 'Neue Bausteine/Designs in JSON eintragen und neu exportieren. Manuelles Zeilenanfügen wird nicht automatisch verdrahtet.'],
        ['Validierung', 'Excel prüft nur elektrische/physische Grenzen; Rolle, Fraktion, Montage und Besatzungskosten-Konsistenz im Python-Rechner prüfen.'],
        ['Bunkerpreis', 'Leeres Bunkermodul plus separat gewählte Startbesatzung. Infanteriepreise und Besatzungszeilen sind editierbar; Kosten/Masse werden einmal addiert.'],
        ['Aufsatzplatz', 'Gatling, Bunker und Lautsprecher teilen roof. Vollständige Slot-/Fraktions-/Kapazitätsprüfung weiterhin im Python-Rechner.'],
        ['Besatzung', '0 in Excel deaktiviert eine vorhandene Zeile; für neue Typen/Zeilen JSON ändern und neu exportieren. Infanteriemasse ist ein Designwert.'],
        ['Momentaufnahme', 'Python-Werte und Preisfit sind statisch. Änderungen in Excel aktualisieren nur die Formelspalten.'],
        ['Genauigkeit', 'CP/Tier sind unabhängig; 1% Abweichung kann akzeptabel sein. Preisfit ist kein Beweis des CP-Wertes.'],
        ['Branches', 'Game modular: tournament-bot integriert in 9d8a9dec; master nicht zusätzlich gemergt.'],
        ['Merge', 'Engine modular-engine: 484afea27c kombiniert Isometric und Tournament.'],
        ['Referenzbedingungen', catalog['reference_conditions']],
    ]:
        info.append(row)
    info.column_dimensions['A'].width = 26
    info.column_dimensions['B'].width = 125

    params = wb.create_sheet('Parameter')
    for row in [('Parameter', 'Wert'), ('Credits pro CP', catalog['credits_per_cp']),
                ('Geschwindigkeitsexponent', catalog['alpha']), ('Katalogbudget', 50)]:
        params.append(row)
    table(params, 'Parameters')
    parts = wb.create_sheet('Bausteine')
    parts.append([label for label, _ in FIELDS])
    for key, part in catalog['components'].items():
        values = [key]
        for _, field in FIELDS[1:]:
            value = part.get(field, '' if field in ('role', 'armor_type', 'note', 'source', 'allowed', 'equipment', 'weapons', 'ammunition') else 0)
            if isinstance(value, (list, dict)):
                value = json.dumps(value, ensure_ascii=False)
            values.append(value)
        parts.append(values)
    table(parts, 'Components')
    end = parts.max_row
    wb.defined_names.add(DefinedName('ComponentIDs', attr_text=f"'Bausteine'!$A$2:$A${end}"))

    infantry = wb.create_sheet('Infanterie')
    infantry.append(['ID', 'Name', 'Fraktionen', 'Preis', 'Masse kg', 'Transportgewicht', 'Tech'])
    for key, unit in catalog['infantry'].items():
        infantry.append([key, unit['display_name'], ', '.join(unit['factions']), unit['cost'], unit['mass'], unit['weight'], unit['tech']])
    table(infantry, 'Infantry')
    wb.defined_names.add(DefinedName('InfantryIDs', attr_text=f"'Infanterie'!$A$2:$A${infantry.max_row}"))
    crew = wb.create_sheet('Startbesatzung')
    crew.append(['Design', 'Infanterie', 'Anzahl', 'Preis', 'Masse kg', 'Transportgewicht', 'Tech'])
    for design in catalog['designs']:
        for unit in design.get('crew', []):
            r = crew.max_row + 1
            crew.append([design['name'], unit['actor'], unit['count']])
            for column, index in [('D', 4), ('E', 5), ('F', 6), ('G', 7)]:
                lookup_unit = f'VLOOKUP(B{r},Infanterie!$A$2:$G${infantry.max_row},{index},FALSE)'
                crew[f'{column}{r}'] = '=' + (f'C{r}*{lookup_unit}' if column != 'G' else f'IF(C{r}>0,{lookup_unit},0)')
    table(crew, 'Crew')
    choices = DataValidation(type='list', formula1='InfantryIDs')
    choices.showErrorMessage = True
    crew.add_data_validation(choices)
    choices.add(f'B2:B{crew.max_row}')
    counts = DataValidation(type='whole', operator='between', formula1=0, formula2=4)
    counts.showErrorMessage = True
    crew.add_data_validation(counts)
    counts.add(f'C2:C{crew.max_row}')

    designs = wb.create_sheet('Entwuerfe')
    designs.append(['Design', 'Fraktion', 'Chassis', 'Motor', 'Generator', 'Panzerung', 'Traeger', 'Waffe', 'Munition',
                    'Zusatz 1', 'Zusatz 2', 'Zusatz 3', 'Katalog', 'CP', 'Tech', 'Hardware', 'CP-Abzug', 'Preis',
                    'HP', 'Masse kg', 'Last kWe', 'Rest kWm', 'Tempo', 'Armor', 'Grenzen', 'Python Preis', 'Python Tempo', 'Referenzhinweis', 'Fertigungsmodul', 'Fertigungsfaktor', 'Hardware nach Fertigung', 'Aufsatztraeger', 'Aufsatzwaffe', 'Aufsatzmunition', 'Besatzungspreis', 'Besatzungsmasse', 'Besatzungstech'])
    roles = ['chassis', 'drive', 'generator', 'armor', 'carrier', 'weapon', 'ammunition']
    def lookup(cell, field):
        index = next(i for i, (_, f) in enumerate(FIELDS, 1) if f == field)
        return f'IF({cell}="",0,VLOOKUP({cell},Bausteine!$A$2:${col(len(FIELDS))}${end},{index},FALSE))'
    for rownum, design in enumerate(catalog['designs'], 2):
        ids = design['components']
        picks = [next(k for k in ids if catalog['components'][k]['role'] == role) for role in roles]
        equipment = [k for k in ids if catalog['components'][k]['role'] == 'equipment']
        if len(equipment) > 3:
            raise ValueError('Workbook template supports at most 3 equipment slots')
        designs.append([design['name'], design['faction']] + picks + equipment + [''] * (3 - len(equipment)))
        designs[f'AB{rownum}'] = design.get('note', '')
        designs[f'AC{rownum}'] = next((k for k in ids if catalog['components'][k]['role'] == 'manufacturing'), '')
        for c, role in [('AF', 'carrier'), ('AG', 'weapon'), ('AH', 'ammunition')]:
            designs[f'{c}{rownum}'] = design.get('auxiliary_mount', {}).get(role, '')
        r = rownum
        def v(c, field):
            return lookup(f'{c}{r}', field)
        def total(field, letters='CDEFGHIJKL'):
            return '+'.join(v(c, field) for c in list(letters) + ['AC', 'AF', 'AG', 'AH'])
        formulas = {
            'M': total('tier'), 'N': total('cp'),
            'O': 'MAX(' + ','.join(v(c, 'tech') for c in list('CDEFGHIJKL') + ['AC', 'AF', 'AG', 'AH']) + f',AK{r})',
            'P': f"{v('C','cost')}*{v('F','cost_percent')}/100+" + total('cost', 'DEGHIJKL') + f'+AI{r}',
            'Q': f'N{r}*Parameter!$B$2', 'R': f'AE{r}-Q{r}',
            'AD': f'IF(AC{r}="",1,' + v('AC','hardware_percent') + '/100)',
            'AE': f'P{r}*AD{r}',
            'S': f"{v('C','hp')}*{v('F','hp_percent')}/100",
            'T': f"{v('C','mass')}*{v('F','mass_percent')}/100+" + total('mass', 'DEGHIJKL') + f'+AJ{r}',
            'AI': f'SUMIF(Startbesatzung!$A$2:$A${crew.max_row},A{r},Startbesatzung!$D$2:$D${crew.max_row})',
            'AJ': f'SUMIF(Startbesatzung!$A$2:$A${crew.max_row},A{r},Startbesatzung!$E$2:$E${crew.max_row})',
            'AK': 'MAX(0,' + ','.join(f'IF(Startbesatzung!A{s}=A{r},Startbesatzung!G{s},0)' for s in range(2, crew.max_row + 1)) + ')',
            'U': total('electric_kw'),
            'V': f"{v('D','mechanical_kw')}-U{r}/{v('E','efficiency')}",
            'W': f'IF(Y{r}<>"OK","",ROUND(MIN(' + v('C','max_speed') + ',' + v('C','reference_speed') +
                 f'*(V{r}/T{r}/(' + v('C','reference_kw') + '/' + v('C','reference_mass') + '))^Parameter!$B$3),0))',
            'X': v('F','armor_type'),
            'Y': f'IFERROR(IF(OR(AD{r}<=0,AD{r}>1,R{r}<=0,T{r}<=0,V{r}<=0,U{r}>' + v('E','max_electric_kw') +
                 f',T{r}>' + v('C','max_mass') + ',' + v('E','efficiency') + '<=0,' + v('E','efficiency') +
                 '>1),"UNGUELTIG","OK"),"FEHLER")',
        }
        for c, formula in formulas.items():
            designs[f'{c}{r}'] = '=' + formula
        try:
            result = calculate(catalog, design)
            designs[f'Z{r}'], designs[f'AA{r}'] = result['cost'], result['speed']
        except ValueError as error:
            designs[f'Z{r}'] = 'INVALID: ' + str(error)
    dv = DataValidation(type='list', formula1='ComponentIDs')
    dv.errorTitle, dv.error = 'Unbekannter Baustein', 'Bitte eine vorhandene ID wählen.'
    dv.showErrorMessage = True
    designs.add_data_validation(dv)
    dv.add(f'C2:L{designs.max_row}')
    dv.add(f'AC2:AC{designs.max_row}')
    dv.add(f'AF2:AH{designs.max_row}')
    table(designs, 'Designs')

    compare = wb.create_sheet('AltNeu')
    compare.append(['Design', 'Preis alt', 'Hardware nach Fertigung', 'CP-Abzug', 'Preis neu', 'Differenz', 'Differenz %',
                    'HP alt', 'HP neu', 'Tempo alt', 'Tempo neu', 'CP neu', 'Katalog neu', 'Grenzen'])
    for source, design in enumerate(catalog['designs'], 2):
        if not design.get('target'):
            continue
        r = compare.max_row + 1
        target = design['target']
        compare.append([design['name'], target['cost'], f'=Entwuerfe!AE{source}', f'=Entwuerfe!Q{source}',
                        f'=Entwuerfe!R{source}', f'=E{r}-B{r}', f'=IF(B{r}=0,0,F{r}/B{r})',
                        target['hp'], f'=Entwuerfe!S{source}', target['speed'], f'=Entwuerfe!W{source}',
                        f'=Entwuerfe!N{source}', f'=Entwuerfe!M{source}', f'=Entwuerfe!Y{source}'])
        compare[f'G{r}'].number_format = '0.0%'
    table(compare, 'Comparison')

    legacy = wb.create_sheet('SchwereReferenzen')
    legacy.append(['Familie', 'Actor', 'Originalpreis', 'PDL Preis', 'PDL Aufpreis', 'Reflector Preis',
                   'Reflector Aufpreis', 'CP ohne Abwehr', 'CP mit Abwehr', 'Basis HP', 'Basistempo', 'Hinweise'])
    for reference in catalog.get('legacy_families', []):
        r = legacy.max_row + 1
        has_defense = 'pdl_price' in reference
        legacy.append([reference['name'], reference['actor'], reference['price'], reference.get('pdl_price'),
                       f'=D{r}-C{r}' if has_defense else None, reference.get('reflector_price'),
                       f'=F{r}-C{r}' if has_defense else None, reference['cp'], reference['cp'] + 1 if has_defense else None,
                       reference['hp'], reference['speed'], 'NUR ORIGINAL, noch kein Baukastendesign. ' + reference['note']])
    table(legacy, 'HeavyReferences')

    fit = wb.create_sheet('Preisfit')
    fit.append(['CP-Wert', 'Fertigungsfaktor', 'Upgrade Hardware', 'PDL Paket', 'Reflector Paket', 'RMSE Credits', 'Max Fehler Credits', 'Status'])
    residuals = wb.create_sheet('FitAbweichungen')
    residuals.append(['CP-Wert', 'Design', 'Preis alt', 'Preis Fit', 'Differenz', 'Differenz %'])
    base, targets = family_targets(catalog)
    for cp in CP_CANDIDATES:
        result = scan(base, targets, cp)
        if result is None:
            fit.append([cp, None, None, None, None, None, None, 'Nicht zulässig unter Modellannahmen'])
            continue
        fit.append([cp] + [result[k] for k in ('factor', 'upgrade', 'pdl', 'reflector', 'rmse', 'max_error')] + ['Statische Fit-Momentaufnahme'])
        for (name, target), predicted in zip(targets, result['predictions']):
            residuals.append([cp, name, target, predicted, predicted - target, (predicted-target)/target])
            residuals.cell(residuals.max_row, 6).number_format = '0.0%'
    table(fit, 'PriceFit')
    table(residuals, 'FitResiduals')
    wb.calculation = CalcProperties(calcId=0, fullCalcOnLoad=True, forceFullCalc=True, calcMode='auto')
    # Exclusive create: never destroy spreadsheet-only what-if edits on regeneration.
    with Path(output).open('xb') as stream:
        wb.save(stream)
    return wb


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output', type=Path, default=Path(__file__).with_name('vehicle-calibration-loadouts.xlsx'))
    args = parser.parse_args()
    data = json.loads(Path(__file__).with_name('catalog.json').read_text(encoding='utf-8'))
    build(data, args.output)
    # Round-trip integrity check; openpyxl cannot evaluate Excel formulas.
    book = load_workbook(args.output)
    assert len(book.sheetnames) == 10
    print(args.output.resolve())
