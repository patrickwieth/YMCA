import importlib.util
import json
from pathlib import Path
import tempfile
import unittest
import subprocess
import sys


@unittest.skipUnless(importlib.util.find_spec('openpyxl'), 'Optional Excel dependency not installed')
class ExcelExportTests(unittest.TestCase):
    def test_cli_roundtrip(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'cli.xlsx'
            result = subprocess.run([sys.executable, str(Path(__file__).with_name('export_excel.py')),
                                     '--output', str(path)], capture_output=True, text=True)
            self.assertEqual(result.returncode, 0, result.stderr)
            self.assertTrue(path.exists())

    def test_roundtrip_formulas_and_overwrite_protection(self):
        from openpyxl import load_workbook
        from export_excel import build
        catalog = json.loads(Path(__file__).with_name('catalog.json').read_text())
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'test.xlsx'
            build(catalog, path)
            book = load_workbook(path)
            self.assertEqual(len(book.sheetnames), 13)
            self.assertIn('FahrzeugFamilien', book.sheetnames)
            jeep = next(row for row in book['FahrzeugFamilien'].iter_rows(min_row=2) if row[1].value == 'JEEP')
            self.assertIn('AA_Jeep', jeep[8].value)
            self.assertNotIn('V3', [row[1].value for row in book['FahrzeugFamilien'].iter_rows(min_row=2)])
            self.assertIn('FahrzeugInventar', book.sheetnames)
            self.assertIn('Fahrprofile', book.sheetnames)
            self.assertEqual(book['Entwuerfe']['AQ2'].value, 'wheels-light')
            for c in ('M', 'N', 'O', 'P', 'T', 'U', 'W', 'Y'):
                self.assertIn('AQ2', book['Entwuerfe'][f'{c}2'].value)
            self.assertIn('VLOOKUP', book['Entwuerfe']['AR2'].value)
            self.assertGreater(book['FahrzeugInventar'].max_row, 100)
            self.assertEqual(book['Infanterie'].max_row, 4)
            self.assertIn('VLOOKUP', book['Startbesatzung']['D2'].value)
            self.assertIn('AF2', book['Entwuerfe']['N2'].value)
            self.assertIn('AI2', book['Entwuerfe']['P2'].value)
            mortar_row = next(row for row in book['Entwuerfe'].iter_rows(min_row=2) if row[0].value == 'Overlord Mortar Roof / draft')
            self.assertEqual(mortar_row[31].value, 'mortar-roof-carrier')
            self.assertEqual(mortar_row[25].value, 2650)
            for cell in mortar_row[37:42]:
                self.assertEqual(cell.data_type, 'f')
            gatling_row = next(row for row in book['Entwuerfe'].iter_rows(min_row=2) if row[0].value == 'Overlord Gatling')
            self.assertEqual(gatling_row[31].value, 'gatling-roof-carrier')
            self.assertEqual(book['SchwereReferenzen'].max_row, 7)
            self.assertEqual(book['SchwereReferenzen']['D3'].value, 2800)
            self.assertEqual(book['SchwereReferenzen']['E3'].value, '=D3-C3')
            self.assertEqual(book['Bausteine'].max_row, len(catalog['components']) + 1)
            self.assertEqual(book['Entwuerfe'].max_row, len(catalog['designs']) + 1)
            self.assertEqual(book['Entwuerfe']['Q2'].value, '=N2*Parameter!$B$2')
            first_reference = next(i for i, d in enumerate(catalog['designs'], 2) if d.get('target'))
            self.assertEqual(book['AltNeu']['E2'].value, f'=Entwuerfe!R{first_reference}')
            self.assertEqual(book['Parameter']['B2'].value, 300)
            self.assertEqual(book['Entwuerfe']['R2'].value, '=AE2-Q2')
            self.assertIn('AC2', book['Entwuerfe']['N2'].value)
            mass_row = next(row for row in book['Entwuerfe'].iter_rows(min_row=2) if row[0].value == 'Battlemaster Mass Production PDL')
            self.assertEqual(mass_row[2].value, 'battlemaster')
            self.assertEqual(mass_row[28].value, 'mass-production')
            self.assertEqual(mass_row[25].value, 1205)
            self.assertTrue(book.calculation.fullCalcOnLoad)
            for row in book['Entwuerfe'].iter_rows(min_row=2):
                for cell in row[12:25]:
                    self.assertEqual(cell.data_type, 'f')
            self.assertEqual(book['Preisfit'].max_row, 10)
            original = path.read_bytes()
            with self.assertRaises(FileExistsError):
                build(catalog, path)
            self.assertEqual(path.read_bytes(), original)


if __name__ == '__main__':
    unittest.main()
